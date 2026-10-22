"""Carga los datos del Excel "ROTULOS RECIBO" a la base de Recibo BOS.

Uso:
    python importar_excel.py "ROTULOS RECIBO V3.0.xlsm"              # portafolio, maestro y personal
    python importar_excel.py "ROTULOS RECIBO V3.0.xlsm" --historico  # además el Consolidado diario
"""
import sys
import unicodedata
from datetime import date, datetime

import openpyxl

import db
from logica import a_entero, calcular

COLUMNAS_PORTAFOLIO = {
    "MATERIAL": "material",
    "TEXTO BREVE DE MATERIAL": "descripcion",
    "UMP": "ump",
    "UMS": "ums",
    "CUADRANTE": "cuadrante",
    "UBICACION": "ubicacion",
    "TENDIDOS": "tendidos",
    "PISOS": "pisos",
    "CAJAS X ESTIBAS": "cajas_estiba",
    "VIDA UTIL CEDI": "vida_util_cedi",
    "VIDA UTIL TIENDA": "vida_util_tienda",
    "CATEGORIA": "categoria",
    "EAN": "ean",
}


def _normalizar(texto):
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(t.upper().split())


def _texto(v):
    if v is None:
        return None
    t = str(v).strip()
    return t or None


def _fecha(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _hoja(wb, nombre):
    for ws in wb.worksheets:
        if _normalizar(ws.title) == _normalizar(nombre):
            return ws
    return None


def leer_portafolio(ws):
    """Devuelve la lista de materiales de una hoja con los encabezados del PORTAFOLIO."""
    filas = ws.iter_rows(values_only=True)
    encabezado = [_normalizar(c) for c in next(filas)]
    indices = {campo: encabezado.index(col) for col, campo in COLUMNAS_PORTAFOLIO.items() if col in encabezado}
    if "material" not in indices or "descripcion" not in indices:
        raise ValueError("La hoja no tiene las columnas MATERIAL y Texto breve de material")

    materiales, vistos = [], set()
    for fila in filas:
        def valor(campo):
            i = indices.get(campo)
            return fila[i] if i is not None and i < len(fila) else None

        material = a_entero(valor("material"))
        descripcion = _texto(valor("descripcion"))
        if not material or not descripcion or material in vistos:
            continue  # igual que BUSCARV: si el material está repetido gana la primera fila
        vistos.add(material)
        tendidos, pisos = a_entero(valor("tendidos")), a_entero(valor("pisos"))
        cajas = a_entero(valor("cajas_estiba")) or (tendidos * pisos if tendidos and pisos else None)
        materiales.append({
            "material": material,
            "descripcion": descripcion,
            "ump": _texto(valor("ump")),
            "ums": _texto(valor("ums")),
            "cuadrante": _texto(valor("cuadrante")),
            "ubicacion": _texto(valor("ubicacion")),
            "tendidos": tendidos,
            "pisos": pisos,
            "cajas_estiba": cajas,
            "vida_util_cedi": a_entero(valor("vida_util_cedi")),
            "vida_util_tienda": a_entero(valor("vida_util_tienda")),
            "categoria": _texto(valor("categoria")),
            "ean": _texto(valor("ean")),
        })
    return materiales


def guardar_materiales(con, materiales, sobrescribir=True):
    campos = list(materiales[0].keys()) if materiales else []
    if not campos:
        return 0
    actualizar = ", ".join(f"{c}=excluded.{c}" for c in campos if c != "material")
    conflicto = (f"DO UPDATE SET {actualizar}, actualizado=datetime('now','localtime')"
                 if sobrescribir else "DO NOTHING")
    con.executemany(
        f"INSERT INTO materiales ({', '.join(campos)}) VALUES ({', '.join(':' + c for c in campos)}) "
        f"ON CONFLICT(material) {conflicto}",
        materiales,
    )
    return len(materiales)


def importar_maestro(con, ws):
    filas = ws.iter_rows(values_only=True)
    encabezado = [_normalizar(c) for c in next(filas)]
    i = {n: encabezado.index(n) for n in ("MATERIAL", "TEXTO BREVE DE MATERIAL", "CODIGO EAN/UPC", "UMS", "UMP", "CATEGORIA")
         if n in encabezado}
    registros = []
    for f in filas:
        material = a_entero(f[i["MATERIAL"]])
        if not material:
            continue
        registros.append((material, _texto(f[i["TEXTO BREVE DE MATERIAL"]]), _texto(f[i["CODIGO EAN/UPC"]]),
                          _texto(f[i["UMS"]]), _texto(f[i["UMP"]]), _texto(f[i["CATEGORIA"]])))
    con.executemany("INSERT OR REPLACE INTO maestro VALUES (?,?,?,?,?,?)", registros)
    return len(registros)


def importar_personal(con, ws):
    nombres = [_texto(ws[f"C{fila}"].value) for fila in range(28, 48)]
    nombres = [n for n in nombres if n]
    con.executemany("INSERT OR IGNORE INTO personal (nombre) VALUES (?)", [(n,) for n in nombres])
    return len(nombres)


def importar_consolidado(con, ws):
    portafolio = {r["material"]: r for r in con.execute("SELECT * FROM materiales")}
    registros = []
    for f in ws.iter_rows(min_row=2, max_col=10, values_only=True):
        fecha, material, cantidad = _fecha(f[0]), a_entero(f[2]), a_entero(f[4])
        if not fecha or not material or not cantidad:
            continue
        m = portafolio.get(material, {})
        cajas = m["cajas_estiba"] if m and m["cajas_estiba"] else cantidad
        completas = a_entero(f[8]) or 0
        saldo = cantidad - completas * cajas if a_entero(f[9]) else 0
        vencimiento = _fecha(f[5])
        vida_min = m["vida_util_cedi"] if m else None
        c = calcular(cantidad, cajas, vencimiento, vida_min, hoy=fecha)
        registros.append({
            "fecha_hora": f"{fecha.isoformat()} 00:00:00", "fecha": fecha.isoformat(), "oc": str(a_entero(f[1]) or ""),
            "material": material, "descripcion": _texto(f[3]), "cantidad": cantidad, "cajas_estiba": cajas,
            "vencimiento": vencimiento.isoformat() if vencimiento else None, "lote": _texto(f[6]),
            "auxiliar": _texto(f[7]), "estibas_completas": completas, "saldo": max(saldo, 0),
            "vida_min": vida_min, "dias_vida_llegada": c["dias_vida"], "status_entrada": c["status"],
            "cuadrante": m["cuadrante"] if m else None, "ubicacion": m["ubicacion"] if m else None,
        })
    registros.reverse()  # el Excel inserta lo más nuevo arriba
    if registros:
        campos = list(registros[0].keys())
        con.executemany(
            f"INSERT INTO recibos ({', '.join(campos)}) VALUES ({', '.join(':' + c for c in campos)})", registros)
    return len(registros)


def main(ruta, historico=False):
    db.inicializar()
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    with db.conectar() as con:
        if (ws := _hoja(wb, "PORTAFOLIO")) is not None:
            print("Portafolio:", guardar_materiales(con, leer_portafolio(ws), sobrescribir=False), "materiales")
        if (ws := _hoja(wb, "Mdt")) is not None:
            print("Maestro SAP:", importar_maestro(con, ws), "materiales")
        if (ws := _hoja(wb, "ROTULO RECIBO FEFO-FIFO")) is not None:
            print("Personal:", importar_personal(con, ws), "personas")
        if historico and (ws := _hoja(wb, "Consolidado diario")) is not None:
            print("Consolidado:", importar_consolidado(con, ws), "recibos")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], historico="--historico" in sys.argv)
